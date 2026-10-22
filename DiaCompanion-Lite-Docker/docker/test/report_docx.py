"""Sinh tài liệu kiểm thử (.docx) từ kết quả chạy — hoặc kế hoạch kiểm thử khi chưa chạy."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Arial"
INK = RGBColor(0x1A, 0x1D, 0x23)
MUTED = RGBColor(0x5A, 0x62, 0x70)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
HEAD_FILL = "E6F2F3"
STATUS_FILL = {"ĐẠT": "E3F4EA", "LỖI": "FBE4E2", "BỊ CHẶN": "FFF1DC", "CHƯA CHẠY": "F1F2F4"}

UI_CHECKLIST = [
    ("UI-01", "Bác sĩ: Thêm bệnh nhân", "Form chỉ có mã nghiên cứu, năm sinh, giới, loại/số năm ĐTĐ, HbA1c nền, ghi chú; KHÔNG có họ tên, SĐT, địa chỉ"),
    ("UI-02", "Nạp ảnh vào thẻ mắt (kéo thả hoặc chọn tệp)", "Thẻ hiện ảnh xem trước cỡ lớn và nút \"Xem ảnh & chấm mù →\"; không có form chấm ngay trên thẻ"),
    ("UI-03", "Mở \"Xem ảnh & chấm mù\"", "Màn toàn màn hình: lăn chuột phóng to, giữ chuột kéo, nhấp đúp về 100%, nút +/−/100%, Red-free đều hoạt động"),
    ("UI-04", "Đọc khung Thông tin lâm sàng ở màn chấm mù", "Hiện đúng tuổi, giới, loại và số năm ĐTĐ, HbA1c lượt khám và HbA1c nền, huyết áp, đường huyết, ghi chú vừa nhập"),
    ("UI-05", "Bấm Esc hoặc ×; mở lại rồi chọn độ và \"Lưu phân độ & chạy AI\"", "Esc/× đóng không lưu; khi lưu, thẻ hiện đồng hồ đếm và sau đó hiện kết quả AI + dòng \"Chấm mù: …\""),
    ("UI-06", "Bấm \"Xem kỹ →\" ở thẻ đã có AI", "Trang kết quả AI có khung Thông tin lâm sàng; ảnh gốc, mask, fractal phóng to/kéo thả được"),
    ("UI-07", "Đăng nhập người chấm", "Vào thẳng trang Chấm ảnh độc lập; ảnh lớn (phóng to, kéo thả, red-free) và khung Thông tin lâm sàng ẩn danh (tuổi, giới, ĐTĐ, HbA1c, HA, đường huyết; KHÔNG mã nghiên cứu, ngày khám, ghi chú); không thấy AI; tab Vòng 3 hoạt động"),
    ("UI-08", "Đăng nhập nghiên cứu, mở Kết xuất dữ liệu", "Có bảng tiến độ chấm mù; tải được 2 tệp CSV; tên tệp mang ngày theo giờ máy"),
    ("UI-09", "Mở Hồ sơ bệnh án của bệnh nhân đã khám", "Mỗi mắt hiện \"Chấm mù: …\" và \"Quyết định: …\" tách riêng"),
    ("UI-10", "Thu hẹp cửa sổ (hoặc mở trên máy tính bảng)", "Bố cục xếp dọc, không tràn ngang; các nút vẫn bấm được"),
    ("UI-11", "Bấm nhanh hai lần (hoặc Enter hai lần) \"Lưu phân độ & chạy AI\"", "Chỉ một ảnh được nạp cho mắt đó; không xuất hiện mục \"Ảnh nạp trùng\"; đóng lượt bình thường"),
    ("UI-12", "Nghiên cứu: Chấm nhiều máy, tạo gói vòng 2, tải zip, giải nén ở máy khác (rút mạng), mở index.html", "Trang chấm chạy không cần mạng; không có mã bệnh nhân, chỉ số lâm sàng, kết quả AI; ảnh không còn EXIF; tải lại trang vẫn giữ phân độ đã chấm"),
    ("UI-13", "Xuất kết quả ở máy chấm, nhập ở máy 1 (Chấm nhiều máy > Nhập kết quả); nhập lại lần nữa", "Lần 1 báo số phân độ mới; lần 2 báo \"đã có từ trước\", không ghi trùng; grading.csv có r2_source = pkg:<mã gói>"),
    ("UI-14", "Tạo gói vòng 3 cho người chấm thứ hai sau khi đã nhập vòng 2", "Gói chỉ chứa ảnh vòng 1 và 2 bất đồng; người chấm vòng 2 không được chọn làm người phân xử cho chính ảnh đó"),
]

INSTALLER_CHECKLIST = [
    ("INS-01", "Cài mới trên máy ảo sạch", "Cài xong không lỗi; trang cuối báo kiểm tra hệ thống OK; có tệp initial_credentials.txt"),
    ("INS-02", "Cài 1.3.1 lên máy đang chạy DiaCompanion Lite 1.3.0", "Tóm tắt báo ĐÃ CHUYỂN; dịch vụ DRAIDiagnosisAI/DRAIDiagnosisBackend chạy, không còn dịch vụ DiaCompanion*; database tên DRAIDiagnosis; dữ liệu cũ đủ trong C:\\ProgramData\\DRAIDiagnosis; có backups\\truoc-doi-ten-*; Start Menu chỉ còn DR AI-Diagnosis; tài khoản cũ đăng nhập được"),
    ("INS-03", "Kiểm tra đúng phiên bản sau khi cài", "DRAIDiagnosis.dll chứa ai_raw_score; ai\\model_runners.py chứa foldScores"),
    ("INS-04", "Chạy từng lối tắt Start Menu", "Kiểm tra hệ thống, Sao lưu, Thêm người chấm, Thêm tài khoản nghiên cứu, Cấp lại mật khẩu, Khởi động lại, Mở thư mục log đều chạy được"),
    ("INS-05", "Mở web từ một máy khác trong mạng LAN", "Truy cập được bằng địa chỉ in trong tệp tóm tắt cài đặt"),
    ("INS-06", "Khởi động lại máy chủ", "Dịch vụ bật đúng theo cấu hình đã chọn (tự động hoặc thủ công); web hoạt động lại"),
    ("INS-07", "Một lượt khám đầy đủ trên bản cài", "Tạo hồ sơ → chấm mù → AI → xác nhận → đóng lượt → tải 2 CSV có cột ai_raw_score có số"),
    ("INS-08", "Sao lưu (Start Menu)", "Có tệp DRAIDiagnosis.bak mới trong C:\\ProgramData\\DRAIDiagnosis\\backups"),
]


# ------------------------------------------------------------------ helpers

def _font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color is not None: run.font.color.rgb = color
    return run


def _p(doc_or_cell, text="", size=10, bold=False, color=INK, align=None, space_after=4, italic=False):
    p = doc_or_cell.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold, color, italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align: p.alignment = align
    return p


def _h(doc, text, level=1):
    p = doc.add_heading(level=level)
    _font(p.add_run(text), 14 if level == 1 else 12, True, TEAL if level == 1 else INK)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def _bullets(doc, items):
    for t in items:
        p = doc.add_paragraph(style="List Bullet")
        _font(p.add_run(t), 10)
        p.paragraph_format.space_after = Pt(2)


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _table(doc, headers, rows, widths_cm, size=9, status_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    _repeat_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        _font(cell.paragraphs[0].add_run(h), size, True)
        _shade(cell, HEAD_FILL)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            bold = status_col is not None and i == status_col
            _font(cells[i].paragraphs[0].add_run(str(v)), size, bold)
            if status_col is not None and i == status_col:
                _shade(cells[i], STATUS_FILL.get(str(v), "FFFFFF"))
    # Độ rộng: khai báo ở cả lưới cột, từng ô và tổng bảng (Word và LibreOffice đọc các chỗ khác nhau).
    tblpr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    # Thứ tự phần tử trong tblPr do lược đồ quy định: tblStyle, tblW, jc, …, tblLayout, tblLook.
    style_el = tblpr.find(qn("w:tblStyle"))
    if style_el is not None:
        style_el.addnext(tblw)
    else:
        tblpr.insert(0, tblw)
    look = tblpr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(layout)
    else:
        tblpr.append(layout)
    for i, w in enumerate(widths_cm):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _section(doc, landscape: bool):
    sec = doc.add_section()
    w, h = Cm(21), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = h, w
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = w, h
    return sec


def _footer(section, text):
    p = section.footer.paragraphs[0]
    p.text = ""
    _font(p.add_run(text + "   ·   Trang "), 8, color=MUTED)
    run = p.add_run()
    _font(run, 8, color=MUTED)
    for kind, val in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = val; run._r.append(it)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


# ------------------------------------------------------------------ report

def build_report(path, registry, modules, meta):
    plan = meta.get("plan_only")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for s in (sec,):
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    title = "KẾ HOẠCH VÀ KỊCH BẢN KIỂM THỬ" if plan else "BÁO CÁO KIỂM THỬ"
    _footer(sec, f"DR AI-Diagnosis {meta['app_version']} · {title.capitalize()}")

    _p(doc, "DR AI-Diagnosis · Hệ thống thu thập dữ liệu sàng lọc võng mạc đái tháo đường", 9, color=MUTED, space_after=2)
    _p(doc, title, 20, True, TEAL, space_after=2)
    _p(doc, f"Phiên bản {meta['app_version']}" + (f" · commit {meta['git_commit']}" if meta.get("git_commit") else ""),
       11, color=MUTED, space_after=10)

    c = meta.get("counts", {})
    n = len(registry)
    info = [
        ("Môi trường", meta["environment"]),
        ("Phạm vi tự động", f"{n} test case API end-to-end, {len({x.module for x in registry})} nhóm chức năng"),
        ("Bắt đầu / kết thúc", "—" if plan else f"{meta['started']} → {meta['finished']} ({meta['duration']})"),
        ("Người thực hiện", meta.get("tester") or "……………………………………"),
    ]
    if not plan:
        info.insert(2, ("Kết quả", f"ĐẠT {c.get('ĐẠT', 0)} · LỖI {c.get('LỖI', 0)} · BỊ CHẶN {c.get('BỊ CHẶN', 0)} / {n}"))
    _table(doc, ["Hạng mục", "Nội dung"], info, [4.5, 12.5])

    if not plan:
        ok = c.get("LỖI", 0) == 0 and c.get("BỊ CHẶN", 0) == 0
        verdict = ("KẾT LUẬN TỰ ĐỘNG: ĐẠT — toàn bộ kịch bản API vượt qua. Hoàn tất checklist thủ công (mục 7, 8) trước khi bàn giao."
                   if ok else f"KẾT LUẬN TỰ ĐỘNG: CHƯA ĐẠT — {c.get('LỖI', 0)} lỗi, {c.get('BỊ CHẶN', 0)} bị chặn. Xem mục 6.")
        p = _p(doc, verdict, 11, True, RGBColor(0x1B, 0x7F, 0x5A) if ok else RGBColor(0xC0, 0x36, 0x2C), space_after=8)

    _h(doc, "1. Mục tiêu và phạm vi")
    _p(doc, "Xác nhận bản phần mềm vận hành đúng các quy tắc của nghiên cứu trước khi triển khai tại bệnh viện:")
    _bullets(doc, [
        "Bảo mật và phân quyền: bác sĩ, người chấm độc lập, nghiên cứu viên chỉ truy cập đúng phần việc.",
        "Ẩn danh ngay khi nhập: hồ sơ chỉ có mã nghiên cứu và năm sinh; tệp kết xuất không lộ định danh.",
        "Chấm mù 3 vòng: vòng 1 bắt buộc trước AI và không sửa được; vòng 2 độc lập; vòng 3 chỉ cho ca bất đồng, do người khác chấm.",
        "Phân tích AI và xác nhận lâm sàng: chặn ảnh không đạt, lưu quyết định riêng với chấm mù, khóa chạy lại sau khi xác nhận.",
        "Dữ liệu nghiên cứu: grading.csv và dataset.csv đủ cột, nhãn tham chiếu đúng quy tắc, điểm AI liên tục và mã băm ảnh chính xác.",
    ])
    _p(doc, "Ngoài phạm vi tự động (kiểm tra thủ công ở mục 7, 8): thao tác giao diện và bộ cài Windows.", italic=True, color=MUTED)

    _h(doc, "2. Môi trường và dữ liệu kiểm thử")
    _table(doc, ["Hạng mục", "Mô tả"], [
        ("Hạ tầng", "Docker Compose, project riêng \"dr-ai-diagnosis-test\": CSDL, ảnh, tài khoản tạo mới mỗi lần chạy và xóa sau khi xong. "
                    "Không dùng dữ liệu của bản dev (cổng 8080) hay máy cài production."),
        ("Thành phần", "SQL Server 2022 Express · Backend ASP.NET Core 8 · Dịch vụ AI (3 model, CPU) · Web nginx"),
        ("Tài khoản", "1 bác sĩ (Doctor + Receptionist), 1 nghiên cứu viên (Research), 2 người chấm độc lập (Grader) — sinh bởi dbinit, mật khẩu tạm đổi trong test"),
        ("Dữ liệu", "2 hồ sơ giả lập NC-TEST-001, NC-TEST-002; ảnh đáy mắt TỔNG HỢP (fixtures/fundus_od.jpg, fundus_os.png), không có ảnh bệnh nhân thật"),
        ("Điểm truy cập API", meta.get("api_url") or "http://backend:8080 (mạng nội bộ Docker)"),
    ], [4, 13])

    _h(doc, "3. Phương pháp")
    _bullets(doc, [
        "Kiểm thử hộp đen qua API HTTP, đúng các lời gọi giao diện web thực hiện; chạy tuần tự theo một luồng nghiên cứu hoàn chỉnh.",
        "Mỗi case kiểm tra mã HTTP và nội dung phản hồi; các con số trong tệp CSV được tính lại độc lập để đối chiếu.",
        "ĐẠT: đúng kết quả mong đợi. LỖI: sai kết quả hoặc lỗi hệ thống. BỊ CHẶN: không chạy được vì case trước nó lỗi. CHƯA CHẠY: kế hoạch.",
        "Chạy lại: docker/test/run-tests.ps1 (Windows) hoặc run-tests.sh — mỗi lần sinh một tài liệu mới.",
    ])

    _h(doc, "4. Tổng hợp theo nhóm chức năng")
    rows, tot = [], [0, 0, 0, 0]
    for code, name in modules.items():
        group = [x for x in registry if x.module == code]
        if not group: continue
        k = [len(group), sum(x.status == "ĐẠT" for x in group), sum(x.status == "LỖI" for x in group), sum(x.status == "BỊ CHẶN" for x in group)]
        tot = [a + b for a, b in zip(tot, k)]
        rows.append((code, name, *k))
    rows.append(("", "TỔNG", *tot))
    _table(doc, ["Mã", "Nhóm chức năng", "Số case", "Đạt", "Lỗi", "Bị chặn"], rows, [1.6, 7.8, 1.9, 1.9, 1.9, 1.9])

    sec2 = _section(doc, landscape=True)
    for s in (sec2,):
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin = s.bottom_margin = Cm(1.6)
    _h(doc, "5. Chi tiết test case")
    _table(doc, ["Mã", "Nội dung kiểm thử", "Kết quả mong đợi", "Kết quả thực tế", "Trạng thái", "Giây"],
           [(x.id, x.title, x.expected, x.actual or "", x.status, "" if plan else f"{x.seconds:.1f}") for x in registry],
           [1.7, 6.3, 6.6, 7.4, 2.0, 1.2], size=8, status_col=4)

    sec3 = _section(doc, landscape=False)
    sec3.left_margin = sec3.right_margin = Cm(2)
    _h(doc, "6. Lỗi phát hiện")
    bad = [x for x in registry if x.status in ("LỖI", "BỊ CHẶN")]
    if plan:
        _p(doc, "Điền sau khi chạy.", italic=True, color=MUTED)
    elif not bad:
        _p(doc, "Không phát hiện lỗi trong phạm vi kiểm thử tự động.")
    else:
        _table(doc, ["Mã", "Trạng thái", "Chi tiết", "Mức độ", "Hướng xử lý"],
               [(x.id, x.status, f"{x.title}\n→ {x.actual}", "", "") for x in bad], [1.6, 1.8, 8.6, 2, 3], status_col=1)

    _h(doc, "7. Checklist giao diện (thủ công)")
    _p(doc, "Thực hiện trên môi trường test (http://localhost:8090) bằng các tài khoản trong test-run/secrets, hoặc trên máy cài sau INS-01.",
       color=MUTED, italic=True)
    _table(doc, ["Mã", "Thao tác", "Kết quả mong đợi", "Đạt?", "Ghi chú"],
           [(a, b, c_, "☐", "") for a, b, c_ in UI_CHECKLIST], [1.4, 4.2, 7.6, 1.2, 2.6])

    _h(doc, "8. Checklist bộ cài Windows (thủ công, trên máy ảo/VPS)")
    _table(doc, ["Mã", "Thao tác", "Kết quả mong đợi", "Đạt?", "Ghi chú"],
           [(a, b, c_, "☐", "") for a, b, c_ in INSTALLER_CHECKLIST], [1.4, 4.2, 7.6, 1.2, 2.6])

    _h(doc, "9. Xác nhận")
    t = _table(doc, ["Người kiểm thử", "Chủ nhiệm đề tài", "Đại diện bệnh viện"],
               [("\n\n\n(ký, ghi rõ họ tên)", "\n\n\n(ký, ghi rõ họ tên)", "\n\n\n(ký, ghi rõ họ tên)"),
                ("Ngày: ……/……/……", "Ngày: ……/……/……", "Ngày: ……/……/……")], [5.6, 5.6, 5.8])
    for row in t.rows[1:]:
        for cell in row.cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Mẫu mặc định của python-docx thiếu w:percent ở w:zoom (Word vẫn mở được, nhưng sai chuẩn).
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.save(path)
    return path
